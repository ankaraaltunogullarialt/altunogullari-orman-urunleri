# stok/views.py
from django.shortcuts import render, get_object_or_404
from django.db.models import Q, Sum, Count, Avg
from django.http import JsonResponse, HttpResponse
from .models import (
    Urun, StokHareket, Ihale, UretimEmri, UretimAsama, Tedarikci,
    IhaleSevk, Tasiyici, TasiyiciArac, TasiyiciOdeme,StokDevir,
)
from personel.models import Personel
from siparis.models import Siparis
from finans.models import CariHesap, Banka
from datetime import datetime, timedelta
from decimal import Decimal
import json


# ==================== RAPORLAMA DASHBOARD ====================
def rapor_dashboard(request):
    """Raporlama ana sayfası"""
    context = {
        'title': 'Raporlama Dashboard',
    }
    return render(request, 'stok/rapor_dashboard.html', context)


# ==================== STOK RAPORU ====================
def stok_raporu(request):
    """Gelişmiş stok raporu"""
    urun_tipi = request.GET.get('urun_tipi', '')
    kategori = request.GET.get('kategori', '')
    stok_durumu = request.GET.get('stok_durumu', '')
    tarih_baslangic = request.GET.get('tarih_baslangic', '')
    tarih_bitis = request.GET.get('tarih_bitis', '')
    
    urunler = Urun.objects.all()
    
    if urun_tipi:
        urunler = urunler.filter(urun_tipi=urun_tipi)
    if kategori:
        urunler = urunler.filter(kategori=kategori)
    if stok_durumu == 'kritik':
        urunler = urunler.filter(mevcut_miktar__lt=10, stokta_mi=True)
    elif stok_durumu == 'stokta':
        urunler = urunler.filter(stokta_mi=True)
    elif stok_durumu == 'tukendi':
        urunler = urunler.filter(mevcut_miktar=0)
    
    toplam_miktar = urunler.aggregate(toplam=Sum('mevcut_miktar'))['toplam'] or 0
    toplam_urun = urunler.count()
    kritik_sayisi = urunler.filter(mevcut_miktar__lt=10, stokta_mi=True).count()
    
    hareketler = StokHareket.objects.all()
    if tarih_baslangic and tarih_bitis:
        hareketler = hareketler.filter(tarih__date__gte=tarih_baslangic, tarih__date__lte=tarih_bitis)
    
    toplam_giris = hareketler.filter(hareket_tipi__in=['ihale_giris', 'sahis_alim', 'mamul_giris']).aggregate(toplam=Sum('miktar'))['toplam'] or 0
    toplam_cikis = hareketler.filter(hareket_tipi__in=['siparis_cikis', 'hammadde_cikis', 'mamul_cikis']).aggregate(toplam=Sum('miktar'))['toplam'] or 0
    
    tip_dagilimi = urunler.values('urun_tipi').annotate(
        adet=Count('id'),
        miktar=Sum('mevcut_miktar')
    )
    
    context = {
        'urunler': urunler,
        'toplam_miktar': toplam_miktar,
        'toplam_urun': toplam_urun,
        'kritik_sayisi': kritik_sayisi,
        'toplam_giris': toplam_giris,
        'toplam_cikis': toplam_cikis,
        'giris_cikis_farki': toplam_giris - toplam_cikis,
        'urun_tipi_secenekleri': Urun.URUN_TIPI,
        'kategori_secenekleri': Urun.URUN_KATEGORI,
        'tip_dagilimi': tip_dagilimi,
    }
    return render(request, 'stok/stok_raporu.html', context)


# ==================== ÜRETİM RAPORU ====================
def uretim_raporu(request):
    """Üretim raporu"""
    toplam_uretim = UretimEmri.objects.filter(durum='tamamlandi').count()
    devam_eden = UretimEmri.objects.filter(durum='devam').count()
    planlanan = UretimEmri.objects.filter(durum='planlandi').count()
    
    aylik_uretim = UretimEmri.objects.filter(
        durum='tamamlandi'
    ).extra(
        select={'ay': "strftime('%%Y-%%m', baslangic_tarihi)"}
    ).values('ay').annotate(
        toplam_miktar=Sum('hedef_miktar'),
        adet=Count('id')
    ).order_by('ay')[:12]
    
    context = {
        'toplam_uretim': toplam_uretim,
        'devam_eden': devam_eden,
        'planlanan': planlanan,
        'aylik_uretim': aylik_uretim,
    }
    return render(request, 'stok/uretim_raporu.html', context)


# ==================== FİNANS RAPORU ====================
def finans_raporu(request):
    """Finans raporu"""
    toplam_alacak = CariHesap.objects.aggregate(toplam=Sum('alacak'))['toplam'] or 0
    toplam_borc = CariHesap.objects.aggregate(toplam=Sum('borc'))['toplam'] or 0
    net_bakiye = toplam_alacak - toplam_borc
    
    bankalar = Banka.objects.filter(aktif_mi=True)
    toplam_banka = bankalar.aggregate(toplam=Sum('bakiye'))['toplam'] or 0
    
    context = {
        'toplam_alacak': toplam_alacak,
        'toplam_borc': toplam_borc,
        'net_bakiye': net_bakiye,
        'bankalar': bankalar,
        'toplam_banka': toplam_banka,
    }
    return render(request, 'stok/finans_raporu.html', context)


# ==================== GRAFİK VERİLERİ (JSON) ====================
def grafik_verileri(request):
    """Chart.js için JSON verileri"""
    son_30_gun = datetime.now() - timedelta(days=30)
    
    gunler = []
    girisler = []
    cikislar = []
    
    for i in range(30, -1, -1):
        gun = datetime.now() - timedelta(days=i)
        gun_baslangic = gun.replace(hour=0, minute=0, second=0)
        gun_bitis = gun.replace(hour=23, minute=59, second=59)
        
        gunler.append(gun.strftime('%d.%m'))
        
        giris = StokHareket.objects.filter(
            tarih__range=(gun_baslangic, gun_bitis),
            hareket_tipi__in=['ihale_giris', 'sahis_alim', 'mamul_giris']
        ).aggregate(toplam=Sum('miktar'))['toplam'] or 0
        girisler.append(float(giris))
        
        cikis = StokHareket.objects.filter(
            tarih__range=(gun_baslangic, gun_bitis),
            hareket_tipi__in=['siparis_cikis', 'hammadde_cikis']
        ).aggregate(toplam=Sum('miktar'))['toplam'] or 0
        cikislar.append(float(cikis))
    
    tip_dagilimi = Urun.objects.values('urun_tipi').annotate(
        adet=Count('id'),
        miktar=Sum('mevcut_miktar')
    )
    
    kritik_urunler = Urun.objects.filter(mevcut_miktar__lt=10, stokta_mi=True)
    kritik_listesi = []
    for urun in kritik_urunler[:10]:
        kritik_listesi.append({
            'ad': urun.urun_adi,
            'miktar': urun.mevcut_miktar
        })
    
    return JsonResponse({
        'gunler': gunler,
        'girisler': girisler,
        'cikislar': cikislar,
        'tip_dagilimi': list(tip_dagilimi),
        'kritik_stok': kritik_listesi,
    })


# ==================== HAREKET ANALİZİ ====================
def hareket_analizi(request):
    """Stok hareket analizi (JSON)"""
    gun_sayisi = int(request.GET.get('gun', 30))
    baslangic = datetime.now() - timedelta(days=gun_sayisi)
    
    hareketler = StokHareket.objects.filter(tarih__gte=baslangic)
    
    gunluk_giris = []
    gunluk_cikis = []
    gunler = []
    
    for i in range(gun_sayisi):
        gun = baslangic + timedelta(days=i)
        gunler.append(gun.strftime('%d.%m'))
        
        gun_hareket = hareketler.filter(tarih__date=gun.date())
        giris = gun_hareket.filter(hareket_tipi='ihale_giris').aggregate(toplam=Sum('miktar'))['toplam'] or 0
        cikis = gun_hareket.filter(hareket_tipi='siparis_cikis').aggregate(toplam=Sum('miktar'))['toplam'] or 0
        
        gunluk_giris.append(float(giris))
        gunluk_cikis.append(float(cikis))
    
    populer_urunler = StokHareket.objects.values('urun__urun_adi').annotate(
        toplam=Sum('miktar')
    ).order_by('-toplam')[:10]
    
    return JsonResponse({
        'gunler': gunler,
        'girisler': gunluk_giris,
        'cikislar': gunluk_cikis,
        'populer_urunler': list(populer_urunler),
    })


# ==================== ÜRETİM ANALİZİ ====================
def uretim_analizi(request):
    """Üretim analizi"""
    aylik_uretim = UretimEmri.objects.filter(durum='tamamlandi').extra(
        select={'ay': "strftime('%%Y-%%m', baslangic_tarihi)"}
    ).values('ay').annotate(
        toplam_miktar=Sum('hedef_miktar'),
        adet=Count('id')
    )
    
    try:
        from .models import UretimAsama
        makine_kullanim = UretimAsama.objects.values('makine__ad').annotate(
            toplam_sure=Sum('tahmini_sure'),
            adet=Count('id')
        )
    except:
        makine_kullanim = []
    
    return JsonResponse({
        'aylik_uretim': list(aylik_uretim),
        'makine_kullanim': list(makine_kullanim),
    })


# ==================== İHALE RAPORU ====================
def ihale_raporu(request):
    """İhale raporu - KPI + KİK Bazlı Toplamlar"""
    from django.db.models import Sum, Count
    
    ihaleler = Ihale.objects.all()
    
    # ===== İSTATİSTİKLER =====
    toplam_ihale = ihaleler.count()
    aktif_ihale = ihaleler.filter(durum='devam_ediyor').count()
    tamamlanan_ihale = ihaleler.filter(durum='tamamlandi').count()
    iptal_ihale = ihaleler.filter(durum='iptal').count()
    
    # ===== MİKTAR İSTATİSTİKLERİ =====
    toplam_ihale_miktari = ihaleler.aggregate(toplam=Sum('toplam_ihale_miktari'))['toplam'] or 0
    toplam_kalan_miktar = ihaleler.aggregate(toplam=Sum('kalan_miktar'))['toplam'] or 0
    toplam_ihale_adet = ihaleler.aggregate(toplam=Sum('toplam_adet'))['toplam'] or 0
    toplam_kalan_adet = ihaleler.aggregate(toplam=Sum('kalan_adet'))['toplam'] or 0
    
    # ===== DEVAM EDEN İHALELER (KİK BAZLI) =====
    # Mantık: Bir KİK No altında en az bir parti devam ediyorsa,
    # o KİK'in TÜM partileri "devam eden" sayılır.
    
    # 1. Hangi KİK No'larda devam eden parti var?
    devam_eden_kikler = ihaleler.filter(
        durum='devam_ediyor'
    ).exclude(
        kik_no__isnull=True
    ).exclude(
        kik_no=''
    ).values_list('kik_no', flat=True).distinct()
    
    # 2. Bu KİK'lerin TÜM partilerini "devam eden" olarak say
    devam_eden_ihaleler = ihaleler.filter(kik_no__in=devam_eden_kikler)
    
    devam_eden_miktar = devam_eden_ihaleler.aggregate(
        toplam=Sum('toplam_ihale_miktari')
    )['toplam'] or 0
    
    devam_eden_kalan = devam_eden_ihaleler.aggregate(
        toplam=Sum('kalan_miktar')
    )['toplam'] or 0
    
    # 3. KİK No'su olmayan ihaleler için ayrı hesapla
    kik_siz_devam = ihaleler.filter(
        durum='devam_ediyor',
        kik_no__isnull=True
    )
    kik_siz_miktar = kik_siz_devam.aggregate(toplam=Sum('toplam_ihale_miktari'))['toplam'] or 0
    kik_siz_kalan = kik_siz_devam.aggregate(toplam=Sum('kalan_miktar'))['toplam'] or 0
    
    # 4. Toplam
    devam_eden_miktar += kik_siz_miktar
    devam_eden_kalan += kik_siz_kalan
    
    # ===== GELEN MİKTAR (Sevklerden) =====
    toplam_gelen_miktar = 0
    for ihale in ihaleler:
        toplam_gelen_miktar += ihale.toplam_gelen_miktar
    
    # ===== TUTAR İSTATİSTİKLERİ =====
    toplam_ihale_tutari = ihaleler.aggregate(toplam=Sum('toplam_tutar'))['toplam'] or 0
    
    # ===== ÖDEME DURUMU =====
    odendiler = ihaleler.filter(odeme_durumu='tamamen_odendi').count()
    kismi_odenmis = ihaleler.filter(odeme_durumu='kismi_odendi').count()
    odenmemis = ihaleler.filter(odeme_durumu='odenmedi').count()
        # Ödeme durumu miktarları (YENİ)
    odendiler_miktar = ihaleler.filter(odeme_durumu='tamamen_odendi').aggregate(t=Sum('toplam_ihale_miktari'))['t'] or 0
    kismi_odenmis_miktar = ihaleler.filter(odeme_durumu='kismi_odendi').aggregate(t=Sum('toplam_ihale_miktari'))['t'] or 0
    odenmemis_miktar = ihaleler.filter(odeme_durumu='odenmedi').aggregate(t=Sum('toplam_ihale_miktari'))['t'] or 0

    # Ödeme durumu kalan miktarları (YENİ)
    odendiler_kalan = ihaleler.filter(odeme_durumu='tamamen_odendi').aggregate(t=Sum('kalan_miktar'))['t'] or 0
    kismi_odenmis_kalan = ihaleler.filter(odeme_durumu='kismi_odendi').aggregate(t=Sum('kalan_miktar'))['t'] or 0
    odenmemis_kalan = ihaleler.filter(odeme_durumu='odenmedi').aggregate(t=Sum('kalan_miktar'))['t'] or 0
    
    # ===== SON 5 İHALE =====
    son_ihaleler = ihaleler.order_by('-ihale_tarihi')[:5]

    # ===== TAMAMEN ÖDENMİŞ İHALELERİN KALAN MİKTARI =====
    tamamen_odenmis = ihaleler.filter(odeme_durumu='tamamen_odendi')
    tamamen_odenmis_kalan = tamamen_odenmis.aggregate(toplam=Sum('kalan_miktar'))['toplam'] or 0
    tamamen_odenmis_toplam = tamamen_odenmis.aggregate(toplam=Sum('toplam_ihale_miktari'))['toplam'] or 0
    tamamen_odenmis_adet = tamamen_odenmis.aggregate(toplam=Sum('kalan_adet'))['toplam'] or 0
    tamamen_odenmis_sayisi = tamamen_odenmis.count()

    
    # ===== KİK BAZLI TOPLAMLAR =====
    kik_nolar = ihaleler.exclude(kik_no__isnull=True).exclude(kik_no='').values_list('kik_no', flat=True).distinct()
    kik_sayisi = kik_nolar.count()
    
    if kik_sayisi > 0:
        kik_bazli_ihaleler = ihaleler.filter(kik_no__in=kik_nolar)
        kik_bazli_toplamlar = kik_bazli_ihaleler.aggregate(
            toplam_miktar=Sum('toplam_ihale_miktari'),
            toplam_kalan=Sum('kalan_miktar'),
            toplam_adet=Sum('toplam_adet'),
            toplam_kalan_adet=Sum('kalan_adet'),
            toplam_tutar=Sum('toplam_tutar'),
            kayit_sayisi=Count('id'),
        )
    else:
        kik_bazli_toplamlar = {
            'toplam_miktar': 0, 'toplam_kalan': 0,
            'toplam_adet': 0, 'toplam_kalan_adet': 0,
            'toplam_tutar': 0, 'kayit_sayisi': 0,
        }
    
    # ===== DEVAM EDEN KİK SAYISI =====
    devam_eden_kik_sayisi = devam_eden_kikler.count()
    devam_eden_parti_sayisi = devam_eden_ihaleler.count()
    
    # ===== CONTEXT =====
    context = {
        'ihaleler': ihaleler,
        'toplam_ihale': toplam_ihale,
        'aktif_ihale': aktif_ihale,
        'tamamlanan_ihale': tamamlanan_ihale,
        'iptal_ihale': iptal_ihale,
        'toplam_ihale_miktari': toplam_ihale_miktari,
        'toplam_kalan_miktar': toplam_kalan_miktar,
        'toplam_gelen_miktar': toplam_gelen_miktar,
        'toplam_ihale_tutari': toplam_ihale_tutari,
        'toplam_ihale_adet': toplam_ihale_adet,
        'toplam_kalan_adet': toplam_kalan_adet,
        'tamamen_odenmis_kalan': tamamen_odenmis_kalan,
        'tamamen_odenmis_toplam': tamamen_odenmis_toplam,
        'tamamen_odenmis_adet': tamamen_odenmis_adet,
        'tamamen_odenmis_sayisi': tamamen_odenmis_sayisi,
        'odendiler': odendiler,
        'odendiler_miktar': odendiler_miktar,
        'kismi_odenmis_miktar': kismi_odenmis_miktar,
        'odenmemis_miktar': odenmemis_miktar,
        'odendiler_kalan': odendiler_kalan,
        'kismi_odenmis_kalan': kismi_odenmis_kalan,
        'odenmemis_kalan': odenmemis_kalan,
        'kismi_odenmis': kismi_odenmis,
        'odenmemis': odenmemis,
        'son_ihaleler': son_ihaleler,
        
        # ===== DEVAM EDEN (KİK BAZLI) =====
        'devam_eden_miktar': devam_eden_miktar,
        'devam_eden_kalan': devam_eden_kalan,
        'devam_eden_kik_sayisi': devam_eden_kik_sayisi,
        'devam_eden_parti_sayisi': devam_eden_parti_sayisi,
        
        # ===== KİK BAZLI =====
        'kik_sayisi': kik_sayisi,
        'kik_bazli_toplam_miktar': kik_bazli_toplamlar['toplam_miktar'] or 0,
        'kik_bazli_toplam_kalan': kik_bazli_toplamlar['toplam_kalan'] or 0,
        'kik_bazli_toplam_adet': kik_bazli_toplamlar['toplam_adet'] or 0,
        'kik_bazli_toplam_kalan_adet': kik_bazli_toplamlar['toplam_kalan_adet'] or 0,
        'kik_bazli_toplam_tutar': kik_bazli_toplamlar['toplam_tutar'] or 0,
        'kik_bazli_kayit_sayisi': kik_bazli_toplamlar['kayit_sayisi'] or 0,
    }
    return render(request, 'stok/ihale_raporu.html', context)


# ==================== İHALE DETAY RAPORU ====================
def ihale_detay_raporu(request, ihale_id=None):
    """
    İhale Detay Raporu
    
    - ihale_id verilirse: Tek bir ihalenin detayı
    - kik_no parametresi verilirse: Aynı KİK'in tüm partileri icmal olarak
    - Diğer filtreler: Boy, Tedarikçi, Tarih aralığı
    """
    from django.db.models import Sum, Count, Avg, Q
    from django.utils.dateparse import parse_date
    from decimal import Decimal
    
    # ===== FİLTRE PARAMETRELERİ =====
    kik_no = request.GET.get('kik_no', '').strip()
    boy = request.GET.get('boy', '').strip()
    tedarikci_id = request.GET.get('tedarikci', '').strip()
    baslangic_str = request.GET.get('baslangic', '').strip()
    bitis_str = request.GET.get('bitis', '').strip()
    
    # ===== TEK İHALE DETAYI =====
    tek_ihale = None
    sevkler = []
    hareketler = []
    toplam_gelen = 0
    toplam_gelen_adet = 0
    
    if ihale_id:
        tek_ihale = get_object_or_404(Ihale, id=ihale_id)
        sevkler = tek_ihale.sevkler.all().order_by('-sevk_tarihi')
        toplam_gelen = sevkler.aggregate(toplam=Sum('sevk_miktar'))['toplam'] or 0
        toplam_gelen_adet = sevkler.aggregate(toplam=Sum('sevk_adet'))['toplam'] or 0
        hareketler = StokHareket.objects.filter(ihale=tek_ihale).order_by('-tarih')
    
    # ===== KİK İCMAL (Aynı KİK'in tüm partileri) =====
    icmal_ihaleler = []
    icmal_toplam = {}
    
    if kik_no:
        icmal_qs = Ihale.objects.filter(kik_no=kik_no).select_related('tedarikci').order_by('parti_no')
        
        if tedarikci_id:
            icmal_qs = icmal_qs.filter(tedarikci_id=tedarikci_id)
        
        icmal_ihaleler = list(icmal_qs)
        
        if icmal_ihaleler:
            agg = icmal_qs.aggregate(
                toplam_miktar=Sum('toplam_ihale_miktari'),
                toplam_kalan=Sum('kalan_miktar'),
                toplam_adet=Sum('toplam_adet'),
                toplam_kalan_adet=Sum('kalan_adet'),
                toplam_tutar=Sum('toplam_tutar'),
            )
            
            miktar = agg['toplam_miktar'] or Decimal('0')
            tutar = agg['toplam_tutar'] or Decimal('0')
            ortalama = (tutar / miktar) if miktar > 0 else Decimal('0')
            
            icmal_toplam = {
                'parti_sayisi': len(icmal_ihaleler),
                'toplam_miktar': miktar,
                'toplam_kalan': agg['toplam_kalan'] or 0,
                'toplam_adet': agg['toplam_adet'] or 0,
                'toplam_kalan_adet': agg['toplam_kalan_adet'] or 0,
                'toplam_tutar': tutar,
                'ortalama_fiyat': ortalama,
            }
    
    # ===== GENEL LİSTE (FİLTRELİ) =====
    ihaleler = None
    if not ihale_id and not kik_no:
        ihaleler = Ihale.objects.select_related('tedarikci').all().order_by('-alis_tarihi', '-id')
        
        if boy:
            ihaleler = ihaleler.filter(boy__icontains=boy)
        if tedarikci_id:
            ihaleler = ihaleler.filter(tedarikci_id=tedarikci_id)
        if baslangic_str:
            baslangic = parse_date(baslangic_str)
            if baslangic:
                ihaleler = ihaleler.filter(alis_tarihi__gte=baslangic)
        if bitis_str:
            bitis = parse_date(bitis_str)
            if bitis:
                ihaleler = ihaleler.filter(alis_tarihi__lte=bitis)
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        if ihale_id and tek_ihale:
            return _ihale_detay_excel(tek_ihale, sevkler)
        elif kik_no and icmal_ihaleler:
            return _ihale_kik_detay_excel(kik_no, icmal_ihaleler, icmal_toplam)
        elif ihaleler is not None:
            return _ihale_liste_excel(ihaleler)
    
    # ===== BOY SEÇENEKLERİ (dropdown için) =====
    boy_secenekleri = Ihale.objects.exclude(boy__isnull=True).exclude(boy='').values_list('boy', flat=True).distinct().order_by('boy')
    
    # ===== KİK NO SEÇENEKLERİ =====
    kik_secenekleri = Ihale.objects.exclude(kik_no__isnull=True).exclude(kik_no='').values_list('kik_no', flat=True).distinct().order_by('kik_no')
    
    context = {
        # Tek ihale
        'tek_ihale': tek_ihale,
        'sevkler': sevkler,
        'hareketler': hareketler,
        'toplam_gelen': toplam_gelen,
        'toplam_gelen_adet': toplam_gelen_adet,
        
        # İcmal
        'kik_no': kik_no,
        'icmal_ihaleler': icmal_ihaleler,
        'icmal_toplam': icmal_toplam,
        
        # Genel liste
        'ihaleler': ihaleler,
        
        # Filtre değerleri
        'boy': boy,
        'tedarikci_id': tedarikci_id,
        'baslangic': baslangic_str,
        'bitis': bitis_str,
        
        # Dropdown seçenekleri
        'boy_secenekleri': boy_secenekleri,
        'kik_secenekleri': kik_secenekleri,
        'tedarikciler': Tedarikci.objects.filter(aktif_mi=True).order_by('unvan'),
    }
    return render(request, 'stok/ihale_detay_raporu.html', context)


# ==================== İHALE İCMAL RAPORU (DÜZELTİLMİŞ) ====================
def ihale_icmal_raporu(request):
    """
    Tüm ihalelerin icmal raporu
    - Filtreleme: Tedarikçi, Parti No, İstif No, KİK No
    - Toplam satırlı
    - Yazdırma desteği
    """
    from django.db.models import Sum, Count, Q
    from .models import Tedarikci  # BU IMPORT EKLENDİ
    
    # Filtreleme parametreleri
    tedarikci_id = request.GET.get('tedarikci', '')
    parti_no = request.GET.get('parti_no', '')
    istif_no = request.GET.get('istif_no', '')
    kik_no = request.GET.get('kik_no', '')
    
    # Sorgu oluştur
    ihaleler = Ihale.objects.all()
    
    if tedarikci_id:
        ihaleler = ihaleler.filter(tedarikci_id=tedarikci_id)
    if parti_no:
        ihaleler = ihaleler.filter(parti_no__icontains=parti_no)
    if istif_no:
        ihaleler = ihaleler.filter(istif_no__icontains=istif_no)
    if kik_no:
        ihaleler = ihaleler.filter(kik_no__icontains=kik_no)
    
    # Benzersiz ihale numaralarını grupla
    benzersiz_ihaleler = ihaleler.values('sistem_ihale_no').annotate(
        toplam_parti=Count('id'),
        toplam_miktar=Sum('toplam_ihale_miktari'),
        toplam_kalan=Sum('kalan_miktar'),
        toplam_adet=Sum('toplam_adet'),
        toplam_kalan_adet=Sum('kalan_adet'),
        toplam_tutar=Sum('toplam_tutar')
    ).order_by('sistem_ihale_no')
    
    # Detayları topla - 4 ONDALIK
    icmal_listesi = []
    for item in benzersiz_ihaleler:
        ihale_no = item['sistem_ihale_no']
        partiler = ihaleler.filter(sistem_ihale_no=ihale_no)
        ilk = partiler.first()
        
        icmal_listesi.append({
            'sistem_ihale_no': ihale_no,
            'kik_no': ilk.kik_no if ilk else '',
            'tedarikci': ilk.tedarikci.unvan if ilk else '',
            'urun_adi': ilk.urun_adi if ilk else '',
            'toplam_parti': item['toplam_parti'],
            'toplam_miktar': round(item['toplam_miktar'] or 0, 4),      # 4 ONDALIK
            'toplam_kalan': round(item['toplam_kalan'] or 0, 4),        # 4 ONDALIK
            'toplam_adet': item['toplam_adet'] or 0,
            'toplam_kalan_adet': item['toplam_kalan_adet'] or 0,
            'toplam_tutar': round(item['toplam_tutar'] or 0, 2),
            'partiler': partiler,
            'durum': ilk.durum if ilk else 'devam_ediyor',
            'odeme_durumu': ilk.odeme_durumu if ilk else 'odenmedi',
        })
    
    # Genel toplamlar - 4 ONDALIK
    genel_toplam = {
        'toplam_ihale': len(benzersiz_ihaleler),
        'toplam_parti': ihaleler.count(),
        'toplam_miktar': round(ihaleler.aggregate(toplam=Sum('toplam_ihale_miktari'))['toplam'] or 0, 4),
        'toplam_kalan': round(ihaleler.aggregate(toplam=Sum('kalan_miktar'))['toplam'] or 0, 4),
        'toplam_adet': ihaleler.aggregate(toplam=Sum('toplam_adet'))['toplam'] or 0,
        'toplam_kalan_adet': ihaleler.aggregate(toplam=Sum('kalan_adet'))['toplam'] or 0,
        'toplam_tutar': round(ihaleler.aggregate(toplam=Sum('toplam_tutar'))['toplam'] or 0, 2),
    }
    
    context = {
        'icmal_listesi': icmal_listesi,
        'genel_toplam': genel_toplam,
        'tedarikciler': tedarikciler,
        'selected_tedarikci': tedarikci_id,
        'parti_no': parti_no,
        'istif_no': istif_no,
        'kik_no': kik_no,
        'title': 'İhale İcmal Raporu',
    }
    return render(request, 'stok/ihale_icmal_raporu.html', context)

# ==================== STOK HAREKET RAPORU (ÖZEL TARİHLİ) ====================
def stok_hareket_raporu(request):
    """
    Stok hareketleri raporu
    - Özel tarih aralığı
    - Hareket tipi filtresi
    - Tedarikçi filtresi
    - Excel export
    """
    from django.utils.dateparse import parse_date
    from datetime import date
    
    # ===== FİLTRE PARAMETRELERİ =====
    baslangic_str = request.GET.get('baslangic', '').strip()
    bitis_str = request.GET.get('bitis', '').strip()
    hareket_tipi = request.GET.get('hareket_tipi', '').strip()
    tedarikci_id = request.GET.get('tedarikci', '').strip()
    
    # ===== QUERYSET =====
    hareketler = StokHareket.objects.select_related(
        'urun', 'ihale', 'ihale__tedarikci', 'ihale_sevk', 'uretim'
    ).all().order_by('-tarih')
    
    # Tarih filtresi
    if baslangic_str:
        baslangic = parse_date(baslangic_str)
        if baslangic:
            hareketler = hareketler.filter(tarih__date__gte=baslangic)
    
    if bitis_str:
        bitis = parse_date(bitis_str)
        if bitis:
            hareketler = hareketler.filter(tarih__date__lte=bitis)
    
    # Hareket tipi filtresi
    if hareket_tipi:
        hareketler = hareketler.filter(hareket_tipi=hareket_tipi)
    
    # Tedarikçi filtresi
    if tedarikci_id:
        hareketler = hareketler.filter(ihale__tedarikci_id=tedarikci_id)
    
    # ===== TOPLAMLAR =====
    toplamlar = hareketler.aggregate(
        toplam_miktar=Sum('miktar'),
        toplam_tutar=Sum('toplam_tutar'),
        kayit_sayisi=Count('id'),
    )
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        return _stok_hareket_excel(hareketler, request)
    
    # ===== CONTEXT =====
    context = {
        'hareketler': hareketler,
        'toplam_miktar': toplamlar['toplam_miktar'] or 0,
        'toplam_tutar': toplamlar['toplam_tutar'] or 0,
        'kayit_sayisi': toplamlar['kayit_sayisi'] or 0,
        
        # Filtre değerleri (form'da kalsın)
        'baslangic': baslangic_str,
        'bitis': bitis_str,
        'hareket_tipi': hareket_tipi,
        'tedarikci_id': tedarikci_id,
        
        # Dropdown seçenekleri
        'hareket_tipi_secenekleri': StokHareket.HAREKET_TIPI,
        'tedarikciler': Tedarikci.objects.filter(aktif_mi=True).order_by('unvan'),
    }
    return render(request, 'stok/stok_hareket_raporu.html', context)


def _stok_hareket_excel(hareketler, request):
    """Stok hareketleri Excel export (openpyxl ile)"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Stok Hareketleri"
    
    # ===== BAŞLIKLAR =====
    headers = [
        'Ürün',
        'Hareket Tipi',
        'Kaynak / Tedarikçi',
        'Parti No',
        'İstif No',
        'Alış Tarihi',
        'Miktar (m³)',
        'Birim Fiyat',
        'Toplam Tutar',
        'Kalan Miktar (m³)',
        'Kalan Adet',
        'Tarih',
    ]
    
    # Başlık stili
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    # ===== VERİLER =====
    for row, h in enumerate(hareketler, 2):
        # Tedarikçi / Kaynak
        if h.ihale and h.ihale.tedarikci:
            kaynak = h.ihale.tedarikci.unvan
        elif h.uretim:
            kaynak = f"Üretim: {h.uretim.uretim_no}"
        else:
            kaynak = h.get_hareket_tipi_display()
        
        # Alış tarihi
        alis_tarihi = ''
        if h.ihale and h.ihale.alis_tarihi:
            alis_tarihi = h.ihale.alis_tarihi.strftime('%d.%m.%Y')
        
        # Kalan miktar/adet
        kalan_miktar = ''
        kalan_adet = ''
        if h.ihale:
            kalan_miktar = float(h.ihale.kalan_miktar) if h.ihale.kalan_miktar else 0
            kalan_adet = h.ihale.kalan_adet or 0
        
        ws.cell(row=row, column=1, value=h.urun.urun_adi if h.urun else '')
        ws.cell(row=row, column=2, value=h.get_hareket_tipi_display())
        ws.cell(row=row, column=3, value=kaynak)
        ws.cell(row=row, column=4, value=h.ihale.parti_no if h.ihale else '')
        ws.cell(row=row, column=5, value=h.ihale.istif_no if h.ihale else '')
        ws.cell(row=row, column=6, value=alis_tarihi)
        ws.cell(row=row, column=7, value=float(h.miktar) if h.miktar else 0)
        ws.cell(row=row, column=8, value=float(h.birim_fiyat) if h.birim_fiyat else 0)
        ws.cell(row=row, column=9, value=float(h.toplam_tutar) if h.toplam_tutar else 0)
        ws.cell(row=row, column=10, value=kalan_miktar)
        ws.cell(row=row, column=11, value=kalan_adet)
        ws.cell(row=row, column=12, value=h.tarih.strftime('%d.%m.%Y %H:%M') if h.tarih else '')
        
        # Border ekle
        for col in range(1, 13):
            ws.cell(row=row, column=col).border = border
    
    # ===== SÜTUN GENİŞLİKLERİ =====
    column_widths = [30, 20, 30, 15, 15, 15, 15, 15, 18, 18, 15, 18]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    # ===== DOSYA ADI =====
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'stok_hareketleri_{tarih_str}.xlsx'
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

# ==================== İHALE RAPORLARI (3'LÜ SET) ====================
from django.db.models import Sum, Count, Avg, F, Q
from decimal import Decimal


# ============================================================
# RAPOR 1: İHALE GENEL RAPORU (Tarih aralığı + tüm bilgiler)
# ============================================================
def ihale_genel_raporu(request):
    """
    Belirli tarih aralığındaki tüm ihaleler
    Filtreler: Başlangıç, Bitiş, Tedarikçi, Durum
    """
    from django.utils.dateparse import parse_date
    
    baslangic_str = request.GET.get('baslangic', '').strip()
    bitis_str = request.GET.get('bitis', '').strip()
    tedarikci_id = request.GET.get('tedarikci', '').strip()
    durum = request.GET.get('durum', '').strip()
    
    ihaleler = Ihale.objects.select_related('tedarikci').all().order_by('-alis_tarihi', '-id')
    
    if baslangic_str:
        baslangic = parse_date(baslangic_str)
        if baslangic:
            ihaleler = ihaleler.filter(alis_tarihi__gte=baslangic)
    
    if bitis_str:
        bitis = parse_date(bitis_str)
        if bitis:
            ihaleler = ihaleler.filter(alis_tarihi__lte=bitis)
    
    if tedarikci_id:
        ihaleler = ihaleler.filter(tedarikci_id=tedarikci_id)
    
    if durum:
        ihaleler = ihaleler.filter(durum=durum)
    
    # ===== TOPLAMLAR (FİLTRELENMİŞ) =====
    toplamlar = ihaleler.aggregate(
        toplam_miktar=Sum('toplam_ihale_miktari'),
        toplam_kalan=Sum('kalan_miktar'),
        toplam_adet=Sum('toplam_adet'),
        toplam_kalan_adet=Sum('kalan_adet'),
        toplam_tutar=Sum('toplam_tutar'),
        kayit_sayisi=Count('id'),
    )
    
    # ===== KİK BAZLI TOPLAMLAR (FİLTRE UYGULANMADAN) =====
    # Filtrelenmiş ihalelerin KİK No'larını al
    kik_nolar = ihaleler.exclude(kik_no__isnull=True).exclude(kik_no='').values_list('kik_no', flat=True).distinct()
    
    # Bu KİK No'lara ait TÜM ihaleleri topla (durum filtresi yok)
    if kik_nolar:
        kik_bazli_ihaleler = Ihale.objects.filter(kik_no__in=kik_nolar)
        kik_bazli_toplamlar = kik_bazli_ihaleler.aggregate(
            toplam_miktar=Sum('toplam_ihale_miktari'),
            toplam_kalan=Sum('kalan_miktar'),
            toplam_adet=Sum('toplam_adet'),
            toplam_kalan_adet=Sum('kalan_adet'),
            toplam_tutar=Sum('toplam_tutar'),
            kayit_sayisi=Count('id'),
        )
    else:
        kik_bazli_toplamlar = {
            'toplam_miktar': 0, 'toplam_kalan': 0,
            'toplam_adet': 0, 'toplam_kalan_adet': 0,
            'toplam_tutar': 0, 'kayit_sayisi': 0,
        }
    
    # ===== ORTALAMA BİRİM FİYAT =====
    # Toplam Tutar / Toplam Miktar
    toplam_tutar = toplamlar['toplam_tutar'] or Decimal('0')
    toplam_miktar = toplamlar['toplam_miktar'] or Decimal('0')
    ortalama_fiyat = (toplam_tutar / toplam_miktar) if toplam_miktar > 0 else Decimal('0')
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        return _ihale_genel_excel(ihaleler, toplamlar, ortalama_fiyat)
    
    context = {
        'ihaleler': ihaleler,
        'toplam_miktar': toplamlar['toplam_miktar'] or 0,
        'toplam_kalan': toplamlar['toplam_kalan'] or 0,
        'toplam_adet': toplamlar['toplam_adet'] or 0,
        'toplam_kalan_adet': toplamlar['toplam_kalan_adet'] or 0,
        'toplam_tutar': toplam_tutar,
        'ortalama_fiyat': ortalama_fiyat,
        'kayit_sayisi': toplamlar['kayit_sayisi'] or 0,
        
        'baslangic': baslangic_str,
        'bitis': bitis_str,
        'tedarikci_id': tedarikci_id,
        'durum': durum,
        
        'tedarikciler': Tedarikci.objects.filter(aktif_mi=True).order_by('unvan'),
        'durum_secenekleri': Ihale.DURUM_CHOICES,
    }
    return render(request, 'stok/ihale_genel_raporu.html', context)


# ============================================================
# RAPOR 2: KİK İCMAL RAPORU (Filtre yoksa tümü)
# ============================================================
def ihale_kik_icmal_raporu(request):
    """
    KİK İcmal Raporu
    
    Filtreler:
    - KİK No (opsiyonel)
    - Tedarikçi (opsiyonel)
    
    Hiçbiri seçilmezse → TÜM ihaleler gösterilir
    """
    from django.db.models import Sum, Count
    from decimal import Decimal
    
    kik_no = request.GET.get('kik_no', '').strip()
    tedarikci_id = request.GET.get('tedarikci', '').strip()
    
    # ===== DROPDOWN SEÇENEKLERİ =====
    kik_nolar = Ihale.objects.exclude(
        kik_no__isnull=True
    ).exclude(
        kik_no=''
    ).values_list('kik_no', flat=True).distinct().order_by('kik_no')
    
    # ===== FİLTRELEME (HER İKİSİ DE OPSİYONEL) =====
    ihaleler = Ihale.objects.select_related('tedarikci').all().order_by('kik_no', 'parti_no')
    
    # KİK No filtresi (varsa)
    if kik_no:
        ihaleler = ihaleler.filter(kik_no=kik_no)
    
    # Tedarikçi filtresi (varsa)
    if tedarikci_id:
        ihaleler = ihaleler.filter(tedarikci_id=tedarikci_id)
    
    # ===== TOPLAMLAR =====
    toplamlar = ihaleler.aggregate(
        toplam_miktar=Sum('toplam_ihale_miktari'),
        toplam_kalan=Sum('kalan_miktar'),
        toplam_adet=Sum('toplam_adet'),
        toplam_kalan_adet=Sum('kalan_adet'),
        toplam_tutar=Sum('toplam_tutar'),
    )
    
    toplam_miktar = toplamlar['toplam_miktar'] or Decimal('0')
    toplam_tutar = toplamlar['toplam_tutar'] or Decimal('0')
    ortalama_fiyat = (toplam_tutar / toplam_miktar) if toplam_miktar > 0 else Decimal('0')
    
    # ===== İCMAL LİSTESİ =====
    icmal_listesi = [{
        'ihale': i,
        'parti_no': i.parti_no,
        'istif_no': i.istif_no,
        'boy': i.boy,
        'tedarikci': i.tedarikci.unvan if i.tedarikci else '-',
        'alis_tarihi': i.alis_tarihi,
        'miktar': i.toplam_ihale_miktari,
        'gelen': i.toplam_gelen_miktar,
        'kalan': i.kalan_miktar,
        'adet': i.toplam_adet,
        'kalan_adet': i.kalan_adet,
        'birim_fiyat': i.birim_fiyat,
        'toplam_tutar': i.toplam_tutar,
        'durum': i.get_durum_display(),
    } for i in ihaleler]
    
    # ===== KİK SAYISI =====
    if kik_no:
        kik_sayisi = 1
    else:
        kik_sayisi = ihaleler.exclude(kik_no__isnull=True).exclude(kik_no='').values('kik_no').distinct().count()
    
    genel_toplam = {
        'toplam_ihale': kik_sayisi,
        'toplam_parti': ihaleler.count(),
        'toplam_miktar': toplam_miktar,
        'toplam_kalan': toplamlar['toplam_kalan'] or 0,
        'toplam_adet': toplamlar['toplam_adet'] or 0,
        'toplam_kalan_adet': toplamlar['toplam_kalan_adet'] or 0,
        'toplam_tutar': toplam_tutar,
        'ortalama_fiyat': ortalama_fiyat,
    }
    
    # ===== FİLTRE BAŞLIĞI =====
    if kik_no and tedarikci_id:
        tedarikci = Tedarikci.objects.filter(id=tedarikci_id).first()
        filtre_baslik = f"KİK No: {kik_no} | Tedarikçi: {tedarikci.unvan if tedarikci else '-'}"
    elif kik_no:
        filtre_baslik = f"KİK No: {kik_no}"
    elif tedarikci_id:
        tedarikci = Tedarikci.objects.filter(id=tedarikci_id).first()
        filtre_baslik = f"Tedarikçi: {tedarikci.unvan if tedarikci else '-'}"
    else:
        filtre_baslik = "TÜM İHALELER"
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        baslik = kik_no if kik_no else 'tum_ihaleler'
        return _ihale_kik_icmal_excel(baslik, icmal_listesi, genel_toplam)
    
    context = {
        'kik_no': kik_no,
        'kik_nolar': kik_nolar,
        'tedarikci_id': tedarikci_id,
        'tedarikciler': Tedarikci.objects.filter(aktif_mi=True).order_by('unvan'),
        'icmal_listesi': icmal_listesi,
        'genel_toplam': genel_toplam,
        'filtre_baslik': filtre_baslik,
    }
    return render(request, 'stok/ihale_kik_icmal_raporu.html', context)


# ============================================================
# RAPOR 3: BOY ANALİZ RAPORU (Hangi boydan kaç m³)
# ============================================================
def ihale_boy_analiz_raporu(request):
    """
    Boy bazlı analiz: Hangi boydan kaç m³ alınmış
    """
    from django.utils.dateparse import parse_date
    
    baslangic_str = request.GET.get('baslangic', '').strip()
    bitis_str = request.GET.get('bitis', '').strip()
    tedarikci_id = request.GET.get('tedarikci', '').strip()
    
    ihaleler = Ihale.objects.exclude(boy__isnull=True).exclude(boy='').all()
    
    if baslangic_str:
        baslangic = parse_date(baslangic_str)
        if baslangic:
            ihaleler = ihaleler.filter(alis_tarihi__gte=baslangic)
    
    if bitis_str:
        bitis = parse_date(bitis_str)
        if bitis:
            ihaleler = ihaleler.filter(alis_tarihi__lte=bitis)
    
    if tedarikci_id:
        ihaleler = ihaleler.filter(tedarikci_id=tedarikci_id)
    
    # ===== BOY BAZLI GRUPLAMA =====
    boy_analizi = ihaleler.values('boy').annotate(
        toplam_miktar=Sum('toplam_ihale_miktari'),
        toplam_kalan=Sum('kalan_miktar'),
        toplam_adet=Sum('toplam_adet'),
        toplam_kalan_adet=Sum('kalan_adet'),
        toplam_tutar=Sum('toplam_tutar'),
        ihale_sayisi=Count('id'),
    ).order_by('-toplam_miktar')
    
    # Ortalama fiyat hesapla (her boy için)
    boy_listesi = []
    for b in boy_analizi:
        miktar = b['toplam_miktar'] or Decimal('0')
        tutar = b['toplam_tutar'] or Decimal('0')
        ortalama = (tutar / miktar) if miktar > 0 else Decimal('0')
        
        boy_listesi.append({
            'boy': b['boy'],
            'toplam_miktar': miktar,
            'toplam_kalan': b['toplam_kalan'] or 0,
            'toplam_adet': b['toplam_adet'] or 0,
            'toplam_kalan_adet': b['toplam_kalan_adet'] or 0,
            'toplam_tutar': tutar,
            'ortalama_fiyat': ortalama,
            'ihale_sayisi': b['ihale_sayisi'],
        })
    
    # ===== GENEL TOPLAM =====
    genel_toplamlar = ihaleler.aggregate(
        toplam_miktar=Sum('toplam_ihale_miktari'),
        toplam_kalan=Sum('kalan_miktar'),
        toplam_adet=Sum('toplam_adet'),
        toplam_tutar=Sum('toplam_tutar'),
    )
    
    genel_miktar = genel_toplamlar['toplam_miktar'] or Decimal('0')
    genel_tutar = genel_toplamlar['toplam_tutar'] or Decimal('0')
    genel_ortalama = (genel_tutar / genel_miktar) if genel_miktar > 0 else Decimal('0')
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        return _ihale_boy_analiz_excel(boy_listesi, genel_toplamlar, genel_ortalama)
    
    context = {
        'boy_listesi': boy_listesi,
        'genel_toplam_miktar': genel_miktar,
        'genel_toplam_kalan': genel_toplamlar['toplam_kalan'] or 0,
        'genel_toplam_adet': genel_toplamlar['toplam_adet'] or 0,
        'genel_toplam_tutar': genel_tutar,
        'genel_ortalama': genel_ortalama,
        'boy_sayisi': len(boy_listesi),
        
        'baslangic': baslangic_str,
        'bitis': bitis_str,
        'tedarikci_id': tedarikci_id,
        'tedarikciler': Tedarikci.objects.filter(aktif_mi=True).order_by('unvan'),
    }
    return render(request, 'stok/ihale_boy_analiz_raporu.html', context)

# ==================== EXCEL EXPORT FONKSİYONLARI ====================
def _ihale_genel_excel(ihaleler, toplamlar, ortalama_fiyat):
    """İhale genel raporu Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "İhale Genel Raporu"
    
    headers = [
        'KİK No', 'Parti No', 'İstif No', 'Boy', 'Tedarikçi', 'Ürün',
        'Alış Tarihi', 'Toplam Miktar (m³)', 'Kalan (m³)',
        'Toplam Adet', 'Kalan Adet', 'Birim Fiyat', 'Toplam Tutar',
        'Durum', 'Ödeme Durumu'
    ]
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    for row, i in enumerate(ihaleler, 2):
        ws.cell(row=row, column=1, value=i.kik_no or '')
        ws.cell(row=row, column=2, value=i.parti_no or '')
        ws.cell(row=row, column=3, value=i.istif_no or '')
        ws.cell(row=row, column=4, value=i.boy or '')
        ws.cell(row=row, column=5, value=i.tedarikci.unvan if i.tedarikci else '')
        ws.cell(row=row, column=6, value=i.urun_adi or '')
        ws.cell(row=row, column=7, value=i.alis_tarihi.strftime('%d.%m.%Y') if i.alis_tarihi else '')
        ws.cell(row=row, column=8, value=float(i.toplam_ihale_miktari or 0))
        ws.cell(row=row, column=9, value=float(i.kalan_miktar or 0))
        ws.cell(row=row, column=10, value=i.toplam_adet or 0)
        ws.cell(row=row, column=11, value=i.kalan_adet or 0)
        ws.cell(row=row, column=12, value=float(i.birim_fiyat or 0))
        ws.cell(row=row, column=13, value=float(i.toplam_tutar or 0))
        ws.cell(row=row, column=14, value=i.get_durum_display())
        ws.cell(row=row, column=15, value=i.get_odeme_durumu_display())
        
        for col in range(1, 16):
            ws.cell(row=row, column=col).border = border
    
    column_widths = [15, 12, 20, 12, 25, 20, 12, 18, 15, 12, 12, 15, 18, 15, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'ihale_genel_raporu_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _ihale_kik_icmal_excel(kik_no, icmal_listesi, genel_toplam):
    """KİK icmal raporu Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"KİK {kik_no}"
    
    # ===== BAŞLIKLAR (KİK No eklendi) =====
    headers = [
        'KİK No',           # ← YENİ
        'Parti No',
        'İstif No',
        'Boy',
        'Tedarikçi',
        'Alış Tarihi',
        'Miktar (m³)',
        'Gelen (m³)',
        'Kalan (m³)',
        'Adet',
        'Kalan Adet',
        'Birim Fiyat',
        'Toplam Tutar',
        'Durum'
    ]
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    row = 2
    for item in icmal_listesi:
        ws.cell(row=row, column=1, value=item['ihale'].kik_no or '')      # ← YENİ
        ws.cell(row=row, column=2, value=item['parti_no'] or '')
        ws.cell(row=row, column=3, value=item['istif_no'] or '')
        ws.cell(row=row, column=4, value=item['boy'] or '')
        ws.cell(row=row, column=5, value=item['tedarikci'])
        ws.cell(row=row, column=6, value=item['alis_tarihi'].strftime('%d.%m.%Y') if item['alis_tarihi'] else '')
        ws.cell(row=row, column=7, value=float(item['miktar'] or 0))
        ws.cell(row=row, column=8, value=float(item['gelen'] or 0))
        ws.cell(row=row, column=9, value=float(item['kalan'] or 0))
        ws.cell(row=row, column=10, value=item['adet'] or 0)
        ws.cell(row=row, column=11, value=item['kalan_adet'] or 0)
        ws.cell(row=row, column=12, value=float(item['birim_fiyat'] or 0))
        ws.cell(row=row, column=13, value=float(item['toplam_tutar'] or 0))
        ws.cell(row=row, column=14, value=item['durum'])
        
        for col in range(1, 15):
            ws.cell(row=row, column=col).border = border
        row += 1
    
    # ===== TOPLAM SATIRI =====
    ws.cell(row=row, column=1, value="TOPLAM")
    ws.cell(row=row, column=7, value=float(genel_toplam['toplam_miktar'] or 0))
    ws.cell(row=row, column=9, value=float(genel_toplam['toplam_kalan'] or 0))
    ws.cell(row=row, column=10, value=genel_toplam['toplam_adet'] or 0)
    ws.cell(row=row, column=11, value=genel_toplam['toplam_kalan_adet'] or 0)
    ws.cell(row=row, column=13, value=float(genel_toplam['toplam_tutar'] or 0))
    
    total_font = Font(bold=True, size=11, color="FFFFFF")
    total_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
    for col in range(1, 15):
        c = ws.cell(row=row, column=col)
        c.font = total_font
        c.fill = total_fill
        c.border = border
    
    # ===== ORTALAMA FİYAT =====
    row += 2
    ws.cell(row=row, column=1, value="Ortalama Alış Fiyatı:")
    ws.cell(row=row, column=1).font = Font(bold=True, size=11)
    ws.cell(row=row, column=2, value=float(genel_toplam['ortalama_fiyat'] or 0))
    ws.cell(row=row, column=2).font = Font(bold=True, size=12, color="0D6EFD")
    ws.cell(row=row, column=3, value="TL/m³")
    
    # ===== SÜTUN GENİŞLİKLERİ =====
    column_widths = [15, 12, 20, 12, 25, 12, 15, 15, 15, 12, 12, 15, 18, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'kik_icmal_{kik_no}_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _ihale_boy_analiz_excel(boy_listesi, genel_toplamlar, genel_ortalama):
    """Boy analiz raporu Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Boy Analizi"
    
    headers = [
        'Boy', 'İhale Sayısı', 'Toplam Miktar (m³)', 'Kalan (m³)',
        'Toplam Adet', 'Kalan Adet', 'Toplam Tutar (TL)', 'Ortalama Fiyat (TL/m³)'
    ]
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    for row, b in enumerate(boy_listesi, 2):
        ws.cell(row=row, column=1, value=b['boy'] or '')
        ws.cell(row=row, column=2, value=b['ihale_sayisi'])
        ws.cell(row=row, column=3, value=float(b['toplam_miktar'] or 0))
        ws.cell(row=row, column=4, value=float(b['toplam_kalan'] or 0))
        ws.cell(row=row, column=5, value=b['toplam_adet'] or 0)
        ws.cell(row=row, column=6, value=b['toplam_kalan_adet'] or 0)
        ws.cell(row=row, column=7, value=float(b['toplam_tutar'] or 0))
        ws.cell(row=row, column=8, value=float(b['ortalama_fiyat'] or 0))
        
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = border
    
    # TOPLAM
    row = len(boy_listesi) + 2
    ws.cell(row=row, column=1, value="GENEL TOPLAM")
    ws.cell(row=row, column=3, value=float(genel_toplamlar['toplam_miktar'] or 0))
    ws.cell(row=row, column=4, value=float(genel_toplamlar['toplam_kalan'] or 0))
    ws.cell(row=row, column=5, value=genel_toplamlar['toplam_adet'] or 0)
    ws.cell(row=row, column=7, value=float(genel_toplamlar['toplam_tutar'] or 0))
    ws.cell(row=row, column=8, value=float(genel_ortalama or 0))
    
    total_font = Font(bold=True, size=11, color="FFFFFF")
    total_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
    for col in range(1, 9):
        c = ws.cell(row=row, column=col)
        c.font = total_font
        c.fill = total_fill
        c.border = border
    
    column_widths = [20, 15, 20, 15, 12, 12, 20, 22]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'boy_analiz_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

def _ihale_detay_excel(ihale, sevkler):
    """İhale detay raporu Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "İhale Detayı"
    
    # ===== BAŞLIKLAR =====
    headers = [
        'KİK No', 'Parti No', 'İstif No', 'Boy', 'Tedarikçi', 'Ürün Adı',
        'Alış Tarihi', 'Toplam Miktar (m³)', 'Gelen Miktar (m³)', 'Kalan Miktar (m³)',
        'Toplam Adet', 'Gelen Adet', 'Kalan Adet', 'Birim Fiyat', 'Toplam Tutar',
        'Durum', 'Ödeme Durumu'
    ]
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    # ===== VERİ =====
    toplam_gelen = sevkler.aggregate(toplam=Sum('sevk_miktar'))['toplam'] or 0
    toplam_gelen_adet = sevkler.aggregate(toplam=Sum('sevk_adet'))['toplam'] or 0
    
    ws.cell(row=2, column=1, value=ihale.kik_no or '')
    ws.cell(row=2, column=2, value=ihale.parti_no or '')
    ws.cell(row=2, column=3, value=ihale.istif_no or '')
    ws.cell(row=2, column=4, value=ihale.boy or '')
    ws.cell(row=2, column=5, value=ihale.tedarikci.unvan if ihale.tedarikci else '')
    ws.cell(row=2, column=6, value=ihale.urun_adi or '')
    ws.cell(row=2, column=7, value=ihale.alis_tarihi.strftime('%d.%m.%Y') if ihale.alis_tarihi else '')
    ws.cell(row=2, column=8, value=float(ihale.toplam_ihale_miktari or 0))
    ws.cell(row=2, column=9, value=float(toplam_gelen))
    ws.cell(row=2, column=10, value=float(ihale.kalan_miktar or 0))
    ws.cell(row=2, column=11, value=ihale.toplam_adet or 0)
    ws.cell(row=2, column=12, value=toplam_gelen_adet)
    ws.cell(row=2, column=13, value=ihale.kalan_adet or 0)
    ws.cell(row=2, column=14, value=float(ihale.birim_fiyat or 0))
    ws.cell(row=2, column=15, value=float(ihale.toplam_tutar or 0))
    ws.cell(row=2, column=16, value=ihale.get_durum_display())
    ws.cell(row=2, column=17, value=ihale.get_odeme_durumu_display())
    
    for col in range(1, 18):
        ws.cell(row=2, column=col).border = border
    
    # ===== SEVKLER SAYFASI =====
    ws2 = wb.create_sheet("Sevkler")
    sevk_headers = ['Sevk Tarihi', 'Sevk Miktarı (m³)', 'Sevk Adeti', 'Kalan Miktar (m³)', 'Kalan Adet']
    
    for col, header in enumerate(sevk_headers, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    for row, s in enumerate(sevkler, 2):
        ws2.cell(row=row, column=1, value=s.sevk_tarihi.strftime('%d.%m.%Y') if s.sevk_tarihi else '')
        ws2.cell(row=row, column=2, value=float(s.sevk_miktar or 0))
        ws2.cell(row=row, column=3, value=s.sevk_adet or 0)
        ws2.cell(row=row, column=4, value=float(s.kalan_miktar or 0))
        ws2.cell(row=row, column=5, value=s.kalan_adet or 0)
        for col in range(1, 6):
            ws2.cell(row=row, column=col).border = border
    
    # ===== SÜTUN GENİŞLİKLERİ =====
    column_widths = [15, 12, 20, 12, 25, 20, 12, 18, 18, 18, 12, 12, 12, 15, 18, 15, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    for i, width in enumerate([15, 18, 12, 18, 12], 1):
        ws2.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'ihale_detay_{ihale.kik_no or ihale.id}_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

def _ihale_kik_detay_excel(kik_no, icmal_ihaleler, icmal_toplam):
    """KİK icmal Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"KİK {kik_no}"
    
    headers = [
        'KİK No', 'Parti No', 'İstif No', 'Boy', 'Tedarikçi',
        'Alış Tarihi', 'Miktar (m³)', 'Kalan (m³)', 'Adet', 'Kalan Adet',
        'Birim Fiyat', 'Toplam Tutar', 'Durum'
    ]
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    row = 2
    for i in icmal_ihaleler:
        ws.cell(row=row, column=1, value=i.kik_no or '')
        ws.cell(row=row, column=2, value=i.parti_no or '')
        ws.cell(row=row, column=3, value=i.istif_no or '')
        ws.cell(row=row, column=4, value=i.boy or '')
        ws.cell(row=row, column=5, value=i.tedarikci.unvan if i.tedarikci else '')
        ws.cell(row=row, column=6, value=i.alis_tarihi.strftime('%d.%m.%Y') if i.alis_tarihi else '')
        ws.cell(row=row, column=7, value=float(i.toplam_ihale_miktari or 0))
        ws.cell(row=row, column=8, value=float(i.kalan_miktar or 0))
        ws.cell(row=row, column=9, value=i.toplam_adet or 0)
        ws.cell(row=row, column=10, value=i.kalan_adet or 0)
        ws.cell(row=row, column=11, value=float(i.birim_fiyat or 0))
        ws.cell(row=row, column=12, value=float(i.toplam_tutar or 0))
        ws.cell(row=row, column=13, value=i.get_durum_display())
        for col in range(1, 14):
            ws.cell(row=row, column=col).border = border
        row += 1
    
    # TOPLAM
    ws.cell(row=row, column=1, value="TOPLAM")
    ws.cell(row=row, column=7, value=float(icmal_toplam['toplam_miktar']))
    ws.cell(row=row, column=8, value=float(icmal_toplam['toplam_kalan']))
    ws.cell(row=row, column=9, value=icmal_toplam['toplam_adet'])
    ws.cell(row=row, column=10, value=icmal_toplam['toplam_kalan_adet'])
    ws.cell(row=row, column=12, value=float(icmal_toplam['toplam_tutar']))
    
    total_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
    for col in range(1, 14):
        c = ws.cell(row=row, column=col)
        c.font = total_font
        c.fill = total_fill
        c.border = border
    
    # ORTALAMA FİYAT
    row += 2
    ws.cell(row=row, column=1, value="Ortalama Alış Fiyatı:")
    ws.cell(row=row, column=1).font = Font(bold=True, size=11)
    ws.cell(row=row, column=2, value=float(icmal_toplam['ortalama_fiyat']))
    ws.cell(row=row, column=2).font = Font(bold=True, size=12, color="0D6EFD")
    ws.cell(row=row, column=3, value="TL/m³")
    
    column_widths = [15, 12, 20, 12, 25, 12, 15, 15, 12, 12, 15, 18, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'kik_detay_{kik_no}_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _ihale_liste_excel(ihaleler):
    """İhale listesi Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "İhaleler"
    
    headers = ['KİK No', 'Parti No', 'Boy', 'Tedarikçi', 'Ürün', 'Alış Tarihi',
               'Miktar (m³)', 'Kalan (m³)', 'Adet', 'Birim Fiyat', 'Toplam Tutar', 'Durum']
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    
    for row, i in enumerate(ihaleler, 2):
        ws.cell(row=row, column=1, value=i.kik_no or '')
        ws.cell(row=row, column=2, value=i.parti_no or '')
        ws.cell(row=row, column=3, value=i.boy or '')
        ws.cell(row=row, column=4, value=i.tedarikci.unvan if i.tedarikci else '')
        ws.cell(row=row, column=5, value=i.urun_adi or '')
        ws.cell(row=row, column=6, value=i.alis_tarihi.strftime('%d.%m.%Y') if i.alis_tarihi else '')
        ws.cell(row=row, column=7, value=float(i.toplam_ihale_miktari or 0))
        ws.cell(row=row, column=8, value=float(i.kalan_miktar or 0))
        ws.cell(row=row, column=9, value=i.toplam_adet or 0)
        ws.cell(row=row, column=10, value=float(i.birim_fiyat or 0))
        ws.cell(row=row, column=11, value=float(i.toplam_tutar or 0))
        ws.cell(row=row, column=12, value=i.get_durum_display())
        for col in range(1, 13):
            ws.cell(row=row, column=col).border = border
    
    for i, width in enumerate([15, 12, 15, 25, 20, 12, 15, 15, 10, 15, 18, 15], 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="ihale_listesi_{tarih_str}.xlsx"'
    wb.save(response)
    return response

# ==================== TAŞIYICI RAPORU ====================
def tasiyici_raporu(request):
    """
    Taşıyıcı bazlı sevk raporu
    - Kim kaça taşımış
    - Nereden nereye
    - Fatura / ödeme durumu
    """
    from django.db.models import Sum, Count
    from decimal import Decimal
    from django.utils.dateparse import parse_date
    
    # ===== FİLTRELER =====
    tasiyici_id = request.GET.get('tasiyici', '').strip()
    baslangic_str = request.GET.get('baslangic', '').strip()
    bitis_str = request.GET.get('bitis', '').strip()
    odeme_durumu = request.GET.get('odeme_durumu', '').strip()
    nereden = request.GET.get('nereden', '').strip()
    
    # ===== QUERYSET =====
    sevkler = IhaleSevk.objects.select_related(
        'ihale', 'ihale__tedarikci', 'tasiyici', 'tasiyici_arac'
    ).exclude(tasiyici__isnull=True).order_by('-sevk_tarihi')
    
    if tasiyici_id:
        sevkler = sevkler.filter(tasiyici_id=tasiyici_id)
    if baslangic_str:
        baslangic = parse_date(baslangic_str)
        if baslangic:
            sevkler = sevkler.filter(sevk_tarihi__gte=baslangic)
    if bitis_str:
        bitis = parse_date(bitis_str)
        if bitis:
            sevkler = sevkler.filter(sevk_tarihi__lte=bitis)
    if odeme_durumu:
        sevkler = sevkler.filter(tasima_odeme_durumu=odeme_durumu)
    if nereden:
        sevkler = sevkler.filter(nereden__icontains=nereden)
    
    # ===== TOPLAMLAR =====
    toplamlar = sevkler.aggregate(
        toplam_miktar=Sum('sevk_miktar'),
        toplam_tasima=Sum('tasima_toplam_tutar'),
        toplam_odenen=Sum('tasima_odenen_tutar'),
        kayit_sayisi=Count('id'),
    )
    
    toplam_tasima = toplamlar['toplam_tasima'] or Decimal('0')
    toplam_odenen = toplamlar['toplam_odenen'] or Decimal('0')
    kalan_odeme = toplam_tasima - toplam_odenen
    
    # ===== TAŞIYICI BAZLI ÖZET =====
    tasiyici_ozet = {}
    for sevk in sevkler:
        if sevk.tasiyici:
            key = sevk.tasiyici.id
            if key not in tasiyici_ozet:
                tasiyici_ozet[key] = {
                    'tasiyici': sevk.tasiyici,
                    'sevk_sayisi': 0,
                    'toplam_miktar': Decimal('0'),
                    'toplam_tasima': Decimal('0'),
                    'toplam_odenen': Decimal('0'),
                    'kalan': Decimal('0'),
                }
            tasiyici_ozet[key]['sevk_sayisi'] += 1
            tasiyici_ozet[key]['toplam_miktar'] += sevk.sevk_miktar or Decimal('0')
            tasiyici_ozet[key]['toplam_tasima'] += sevk.tasima_toplam_tutar or Decimal('0')
            tasiyici_ozet[key]['toplam_odenen'] += sevk.tasima_odenen_tutar or Decimal('0')
    
    for k, v in tasiyici_ozet.items():
        v['kalan'] = v['toplam_tasima'] - v['toplam_odenen']
    
    tasiyici_listesi = sorted(
        tasiyici_ozet.values(),
        key=lambda x: x['toplam_miktar'],
        reverse=True
    )
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        return _tasiyici_excel(sevkler, tasiyici_listesi, toplamlar)
    
    # ===== NEREDEN SEÇENEKLERİ (dropdown için) =====
    nereden_secenekleri = IhaleSevk.objects.exclude(
        nereden__isnull=True
    ).exclude(
        nereden=''
    ).values_list('nereden', flat=True).distinct().order_by('nereden')
    
    # ===== CONTEXT =====
    context = {
        'sevkler': sevkler,
        'tasiyici_listesi': tasiyici_listesi,
        'toplam_miktar': toplamlar['toplam_miktar'] or 0,
        'toplam_tasima': toplam_tasima,
        'toplam_odenen': toplam_odenen,
        'kalan_odeme': kalan_odeme,
        'kayit_sayisi': toplamlar['kayit_sayisi'] or 0,
        
        # Filtre değerleri
        'tasiyici_id': tasiyici_id,
        'baslangic': baslangic_str,
        'bitis': bitis_str,
        'odeme_durumu': odeme_durumu,
        'nereden': nereden,
        
        # Dropdown seçenekleri
        'tasiyicilar': Tasiyici.objects.filter(aktif_mi=True).order_by('ad'),
        'odeme_durumlari': IhaleSevk.ODEME_DURUMU_CHOICES,
        'nereden_secenekleri': nereden_secenekleri,
    }
    return render(request, 'stok/tasiyici_raporu.html', context)


def _tasiyici_excel(sevkler, tasiyici_listesi, toplamlar):
    """Taşıyıcı raporu Excel export"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    
    # ===== SAYFA 1: SEVK DETAYLARI =====
    ws1 = wb.active
    ws1.title = "Sevk Detayları"
    
    headers = [
        'Sevk Tarihi', 'İhale No', 'Parti No', 'Taşıyıcı', 'Araç Plaka', 'Şoför',
        'İrsaliye No', 'Nereden', 'Nereye',
        'Sevk Miktarı (m³)', 'Sevk Adeti',
        'Fatura No', 'Fatura Tarihi',
        'Taşıma Birim Fiyat (TL/m³)', 'Taşıma Toplam (TL)',
        'Ödenen (TL)', 'Ödeme Durumu', 'Ödeme Tarihi'
    ]
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    
    for row, s in enumerate(sevkler, 2):
        ws1.cell(row=row, column=1, value=s.sevk_tarihi.strftime('%d.%m.%Y') if s.sevk_tarihi else '')
        ws1.cell(row=row, column=2, value=s.ihale.sistem_ihale_no if s.ihale else '')
        ws1.cell(row=row, column=3, value=s.ihale.parti_no if s.ihale else '')
        ws1.cell(row=row, column=4, value=s.tasiyici.ad if s.tasiyici else '')
        ws1.cell(row=row, column=5, value=s.tasiyici_arac.plaka if s.tasiyici_arac else '')
        ws1.cell(row=row, column=6, value=s.tasiyici_arac.sofor_adi if s.tasiyici_arac else '')
        ws1.cell(row=row, column=7, value=s.irsaliye_no or '')
        ws1.cell(row=row, column=8, value=s.nereden or '')
        ws1.cell(row=row, column=9, value=s.nereye or '')
        ws1.cell(row=row, column=10, value=float(s.sevk_miktar or 0))
        ws1.cell(row=row, column=11, value=s.sevk_adet or 0)
        ws1.cell(row=row, column=12, value=s.fatura_no or '')
        ws1.cell(row=row, column=13, value=s.fatura_tarihi.strftime('%d.%m.%Y') if s.fatura_tarihi else '')
        ws1.cell(row=row, column=14, value=float(s.tasima_birim_fiyat or 0))
        ws1.cell(row=row, column=15, value=float(s.tasima_toplam_tutar or 0))
        ws1.cell(row=row, column=16, value=float(s.tasima_odenen_tutar or 0))
        ws1.cell(row=row, column=17, value=s.get_tasima_odeme_durumu_display())
        ws1.cell(row=row, column=18, value='')

        for col in range(1, 19):
            ws1.cell(row=row, column=col).border = border
    
    # ===== SÜTUN GENİŞLİKLERİ =====
    column_widths = [12, 15, 12, 25, 15, 18, 15, 20, 20, 15, 12, 15, 12, 18, 18, 15, 15, 12]
    for i, width in enumerate(column_widths, 1):
        ws1.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = width
    
    # ===== SAYFA 2: TAŞIYICI ÖZET =====
    ws2 = wb.create_sheet("Taşıyıcı Özet")
    
    headers2 = ['Taşıyıcı', 'Sevk Sayısı', 'Toplam Miktar (m³)', 'Toplam Taşıma (TL)', 'Ödenen (TL)', 'Kalan (TL)']
    
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    
    for row, t in enumerate(tasiyici_listesi, 2):
        ws2.cell(row=row, column=1, value=t['tasiyici'].ad)
        ws2.cell(row=row, column=2, value=t['sevk_sayisi'])
        ws2.cell(row=row, column=3, value=float(t['toplam_miktar']))
        ws2.cell(row=row, column=4, value=float(t['toplam_tasima']))
        ws2.cell(row=row, column=5, value=float(t['toplam_odenen']))
        ws2.cell(row=row, column=6, value=float(t['kalan']))
        
        for col in range(1, 7):
            ws2.cell(row=row, column=col).border = border
    
    for i, width in enumerate([30, 12, 20, 20, 18, 18], 1):
        ws2.column_dimensions[chr(64 + i)].width = width
    
    # ===== TOPLAM SATIRI =====
    row = len(tasiyici_listesi) + 2
    ws2.cell(row=row, column=1, value="TOPLAM")
    ws2.cell(row=row, column=3, value=float(toplamlar['toplam_miktar'] or 0))
    ws2.cell(row=row, column=4, value=float(toplamlar['toplam_tasima'] or 0))
    ws2.cell(row=row, column=5, value=float(toplamlar['toplam_odenen'] or 0))
    
    total_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
    for col in range(1, 7):
        c = ws2.cell(row=row, column=col)
        c.font = total_font
        c.fill = total_fill
        c.border = border
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'tasiyici_raporu_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

# ==================== VERİTABANI YEDEKLEME ====================
import os
import io
import json
import shutil
from datetime import datetime
from django.conf import settings
from django.http import FileResponse, Http404, HttpResponse
from django.contrib import messages
from django.shortcuts import redirect, render
from django.core.management import call_command


# Yedek klasörü
YEDEK_KLASORU = os.path.join(settings.BASE_DIR, 'yedekler')


def _yedek_klasoru_olustur():
    """Yedek klasörü yoksa oluştur"""
    if not os.path.exists(YEDEK_KLASORU):
        os.makedirs(YEDEK_KLASORU)


def _veritabani_turu():
    """Aktif veritabanı türünü döndürür: 'sqlite' veya 'postgresql'"""
    engine = settings.DATABASES['default']['ENGINE']
    if 'sqlite' in engine:
        return 'sqlite'
    elif 'postgresql' in engine:
        return 'postgresql'
    return 'unknown'


def yedekleme_sayfasi(request):
    """Veritabanı yedekleme sayfası"""
    _yedek_klasoru_olustur()

    yedekler = []
    if os.path.exists(YEDEK_KLASORU):
        for dosya in os.listdir(YEDEK_KLASORU):
            if dosya.endswith(('.sqlite3', '.zip', '.sql', '.json')):
                dosya_yolu = os.path.join(YEDEK_KLASORU, dosya)
                if os.path.isfile(dosya_yolu):
                    boyut = os.path.getsize(dosya_yolu)
                    tarih = datetime.fromtimestamp(os.path.getmtime(dosya_yolu))
                    yedekler.append({
                        'ad': dosya,
                        'boyut': boyut,
                        'boyut_mb': round(boyut / (1024 * 1024), 2),
                        'tarih': tarih,
                    })

    yedekler.sort(key=lambda x: x['tarih'], reverse=True)

    # Veritabanı bilgisi
    db_turu = _veritabani_turu()
    if db_turu == 'sqlite':
        db_yolu = settings.DATABASES['default']['NAME']
        db_boyut = os.path.getsize(db_yolu) if os.path.exists(str(db_yolu)) else 0
        db_bilgi = f"SQLite ({db_boyut / 1024:.1f} KB)"
    else:
        # Neon PostgreSQL — boyutu sorgula
        try:
            from django.db import connection
            with connection.cursor() as cur:
                cur.execute("SELECT pg_database_size(current_database())")
                db_boyut = cur.fetchone()[0]
            db_bilgi = f"PostgreSQL / Neon ({db_boyut / (1024 * 1024):.2f} MB)"
        except Exception:
            db_bilgi = "PostgreSQL / Neon"
            db_boyut = 0

    context = {
        'yedekler': yedekler,
        'db_bilgi': db_bilgi,
        'db_boyut_mb': round(db_boyut / (1024 * 1024), 2),
        'yedek_sayisi': len(yedekler),
        'toplam_yedek_boyut': round(sum(y['boyut'] for y in yedekler) / (1024 * 1024), 2),
        'db_turu': db_turu,
    }
    return render(request, 'stok/yedekleme.html', context)


def yedek_olustur(request):
    """Yeni yedek oluştur — SQLite ise kopyala, PostgreSQL ise dumpdata ile JSON al"""
    _yedek_klasoru_olustur()

    try:
        tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        db_turu = _veritabani_turu()

        if db_turu == 'sqlite':
            # === SQLite: dosya kopyala ===
            db_yolu = str(settings.DATABASES['default']['NAME'])
            if not os.path.exists(db_yolu):
                messages.error(request, '❌ Veritabanı dosyası bulunamadı!')
                return redirect('stok:yedekleme_sayfasi')

            yedek_ad = f'db_yedek_{tarih_str}.sqlite3'
            yedek_yolu = os.path.join(YEDEK_KLASORU, yedek_ad)
            shutil.copy2(db_yolu, yedek_yolu)

        else:
            # === PostgreSQL / Neon: dumpdata ile JSON al ===
            yedek_ad = f'db_yedek_{tarih_str}.json'
            yedek_yolu = os.path.join(YEDEK_KLASORU, yedek_ad)

            buffer = io.StringIO()
            call_command(
                'dumpdata',
                '--natural-foreign',
                '--natural-primary',
                '--exclude=contenttypes',
                '--exclude=auth.permission',
                '--indent', '2',
                stdout=buffer,
            )
            with open(yedek_yolu, 'w', encoding='utf-8') as f:
                f.write(buffer.getvalue())

        boyut_mb = round(os.path.getsize(yedek_yolu) / (1024 * 1024), 2)
        messages.success(request, f'✅ Yedek oluşturuldu: {yedek_ad} ({boyut_mb} MB)')

    except Exception as e:
        messages.error(request, f'❌ Yedek oluşturma hatası: {str(e)}')

    return redirect('stok:yedekleme_sayfasi')


def yedek_indir(request, yedek_adi):
    """Yedek dosyasını indir"""
    if '..' in yedek_adi or '/' in yedek_adi or '\\' in yedek_adi:
        raise Http404("Geçersiz dosya adı")

    yedek_yolu = os.path.join(YEDEK_KLASORU, yedek_adi)

    if not os.path.exists(yedek_yolu):
        raise Http404("Yedek bulunamadı")

    response = FileResponse(open(yedek_yolu, 'rb'), as_attachment=True)
    response['Content-Disposition'] = f'attachment; filename="{yedek_adi}"'
    return response


def yedek_sil(request, yedek_adi):
    """Yedek dosyasını sil"""
    if '..' in yedek_adi or '/' in yedek_adi or '\\' in yedek_adi:
        messages.error(request, '❌ Geçersiz dosya adı!')
        return redirect('stok:yedekleme_sayfasi')

    yedek_yolu = os.path.join(YEDEK_KLASORU, yedek_adi)

    try:
        if os.path.exists(yedek_yolu):
            os.remove(yedek_yolu)
            messages.success(request, f'✅ Silindi: {yedek_adi}')
        else:
            messages.warning(request, f'⚠️ Dosya bulunamadı: {yedek_adi}')
    except Exception as e:
        messages.error(request, f'❌ Silme hatası: {str(e)}')

    return redirect('stok:yedekleme_sayfasi')


def yedek_temizle(request):
    """Tüm eski yedekleri temizle (son 5 hariç)"""
    _yedek_klasoru_olustur()

    try:
        yedekler = []
        for dosya in os.listdir(YEDEK_KLASORU):
            if dosya.endswith(('.sqlite3', '.json')):
                yol = os.path.join(YEDEK_KLASORU, dosya)
                yedekler.append((yol, os.path.getmtime(yol)))

        yedekler.sort(key=lambda x: x[1], reverse=True)

        silinen = 0
        for yol, _ in yedekler[5:]:
            os.remove(yol)
            silinen += 1

        messages.success(request, f'✅ {silinen} eski yedek silindi')
    except Exception as e:
        messages.error(request, f'❌ Temizleme hatası: {str(e)}')

    return redirect('stok:yedekleme_sayfasi')

def yedek_dosyadan_yukle(request):
    """Bilgisayardan yüklenen JSON dosyasını veritabanına aktar"""
    if request.method != 'POST':
        return redirect('stok:yedekleme_sayfasi')

    yuklenen_dosya = request.FILES.get('yedek_dosyasi')

    if not yuklenen_dosya:
        messages.error(request, '❌ Dosya seçilmedi!')
        return redirect('stok:yedekleme_sayfasi')

    if not yuklenen_dosya.name.endswith('.json'):
        messages.error(request, '❌ Sadece JSON dosyaları yüklenebilir!')
        return redirect('stok:yedekleme_sayfasi')

    try:
        # Yüklenen dosyayı geçici olarak kaydet
        tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        gecici_ad = f'yuklenen_{tarih_str}.json'
        gecici_yol = os.path.join(YEDEK_KLASORU, gecici_ad)

        with open(gecici_yol, 'wb+') as f:
            for chunk in yuklenen_dosya.chunks():
                f.write(chunk)

        # Yüklemeden önce otomatik yedek al
        otomatik_yedek_ad = f'otomatik_dosya_yukleme_oncesi_{tarih_str}.json'
        otomatik_yedek_yolu = os.path.join(YEDEK_KLASORU, otomatik_yedek_ad)

        buffer = io.StringIO()
        call_command(
            'dumpdata',
            '--natural-foreign',
            '--natural-primary',
            '--exclude=contenttypes',
            '--exclude=auth.permission',
            '--indent', '2',
            stdout=buffer,
        )
        with open(otomatik_yedek_yolu, 'w', encoding='utf-8') as f:
            f.write(buffer.getvalue())

        # loaddata ile yükle
        call_command('loaddata', gecici_yol, verbosity=0)

        messages.success(
            request,
            f'✅ Dosya başarıyla yüklendi: {yuklenen_dosya.name} | '
            f'ℹ️ Önceki durum yedeklendi: {otomatik_yedek_ad}'
        )
    except Exception as e:
        messages.error(request, f'❌ Yükleme hatası: {str(e)}')

    return redirect('stok:yedekleme_sayfasi')

def yedek_yukle(request):
    """JSON yedeğinden geri yükle (dikkatli kullanın)"""
    if request.method != 'POST':
        return redirect('stok:yedekleme_sayfasi')

    yedek_adi = request.POST.get('yedek_adi', '').strip()

    if not yedek_adi:
        messages.error(request, '❌ Yedek seçilmedi!')
        return redirect('stok:yedekleme_sayfasi')

    # Güvenlik
    if '..' in yedek_adi or '/' in yedek_adi or '\\' in yedek_adi:
        messages.error(request, '❌ Geçersiz dosya adı!')
        return redirect('stok:yedekleme_sayfasi')

    yedek_yolu = os.path.join(YEDEK_KLASORU, yedek_adi)

    if not os.path.exists(yedek_yolu):
        messages.error(request, f'❌ Yedek bulunamadı: {yedek_adi}')
        return redirect('stok:yedekleme_sayfasi')

    # Sadece JSON yedeklerini kabul et (PostgreSQL uyumlu)
    if not yedek_adi.endswith('.json'):
        messages.error(request, '❌ Sadece JSON yedekleri geri yüklenebilir!')
        return redirect('stok:yedekleme_sayfasi')

    try:
        # Yüklemeden önce otomatik yedek al
        tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        otomatik_yedek_ad = f'otomatik_geri_yukleme_oncesi_{tarih_str}.json'
        otomatik_yedek_yolu = os.path.join(YEDEK_KLASORU, otomatik_yedek_ad)

        buffer = io.StringIO()
        call_command(
            'dumpdata',
            '--natural-foreign',
            '--natural-primary',
            '--exclude=contenttypes',
            '--exclude=auth.permission',
            '--indent', '2',
            stdout=buffer,
        )
        with open(otomatik_yedek_yolu, 'w', encoding='utf-8') as f:
            f.write(buffer.getvalue())

        # loaddata ile geri yükle
        call_command('loaddata', yedek_yolu, verbosity=0)

        messages.success(
            request,
            f'✅ Geri yükleme başarılı: {yedek_adi} | '
            f'ℹ️ Önceki durum yedeklendi: {otomatik_yedek_ad}'
        )
    except Exception as e:
        messages.error(request, f'❌ Geri yükleme hatası: {str(e)}')

    return redirect('stok:yedekleme_sayfasi')

# ==================== ÜRÜN ANALİZ RAPORU ====================
def ihale_urun_analiz_raporu(request):
    """
    Ürün adı bazlı analiz raporu
    Kategoriler: Göknar, Ladin, Karaçam, Sarıçam
    """
    from django.db.models import Sum, Count, Q
    from decimal import Decimal
    
    # ===== KATEGORİLER =====
    kategoriler = [
        {
            'ad': 'Göknar',
            'sorgu': Q(urun_adi__icontains='GÖKNAR'),
            'renk': 'primary',
            'icon': 'fa-tree',
        },
        {
            'ad': 'Ladin',
            'sorgu': Q(urun_adi__icontains='LADİN'),
            'renk': 'info',
            'icon': 'fa-tree',
        },
        {
            'ad': 'Karaçam',
            'sorgu': Q(urun_adi__icontains='ÇK') | Q(urun_adi__icontains='KARAÇAM'),
            'renk': 'warning',
            'icon': 'fa-tree',
        },
        {
            'ad': 'Sarıçam',
            'sorgu': Q(urun_adi__icontains='ÇS') | Q(urun_adi__icontains='SARIÇAM'),
            'renk': 'success',
            'icon': 'fa-tree',
        },
    ]
    
    # ===== HER KATEGORİ İÇİN HESAPLA =====
    sonuclar = []
    genel_toplam = {
        'ihale_sayisi': 0,
        'toplam_miktar': Decimal('0'),
        'kalan_miktar': Decimal('0'),
        'toplam_tutar': Decimal('0'),
    }
    
    for kat in kategoriler:
        ihaleler = Ihale.objects.filter(kat['sorgu'])
        
        toplamlar = ihaleler.aggregate(
            toplam_miktar=Sum('toplam_ihale_miktari'),
            kalan_miktar=Sum('kalan_miktar'),
            toplam_tutar=Sum('toplam_tutar'),
        )
        
        sonuc = {
            'ad': kat['ad'],
            'renk': kat['renk'],
            'icon': kat['icon'],
            'ihale_sayisi': ihaleler.count(),
            'toplam_miktar': toplamlar['toplam_miktar'] or Decimal('0'),
            'kalan_miktar': toplamlar['kalan_miktar'] or Decimal('0'),
            'toplam_tutar': toplamlar['toplam_tutar'] or Decimal('0'),
        }
        sonuclar.append(sonuc)
        
        # Genel toplama ekle
        genel_toplam['ihale_sayisi'] += sonuc['ihale_sayisi']
        genel_toplam['toplam_miktar'] += sonuc['toplam_miktar']
        genel_toplam['kalan_miktar'] += sonuc['kalan_miktar']
        genel_toplam['toplam_tutar'] += sonuc['toplam_tutar']
    
    # ===== EN ÇOK VE EN AZ =====
    en_cok = max(sonuclar, key=lambda x: x['toplam_miktar']) if sonuclar else None
    en_az = min(sonuclar, key=lambda x: x['toplam_miktar']) if sonuclar else None
    
    # ===== TOPLAM MİKTAR (GENEL) =====
    toplam_ihale = Ihale.objects.count()
    toplam_ihale_miktari = Ihale.objects.aggregate(
        t=Sum('toplam_ihale_miktari')
    )['t'] or Decimal('0')
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        return _ihale_urun_analiz_excel(sonuclar, genel_toplam)
    
    context = {
        'sonuclar': sonuclar,
        'genel_toplam': genel_toplam,
        'en_cok': en_cok,
        'en_az': en_az,
        'toplam_ihale': toplam_ihale,
        'toplam_ihale_miktari': toplam_ihale_miktari,
    }
    return render(request, 'stok/ihale_urun_analiz_raporu.html', context)


def _ihale_urun_analiz_excel(sonuclar, genel_toplam):
    """Ürün analiz raporu Excel export"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Ürün Analizi"
    
    # ===== BAŞLIKLAR =====
    headers = ['Ürün', 'İhale Sayısı', 'Toplam (m³)', 'Kalan (m³)', 'Tutar (TL)']
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    # ===== VERİLER =====
    row = 2
    for s in sonuclar:
        ws.cell(row=row, column=1, value=s['ad'])
        ws.cell(row=row, column=2, value=s['ihale_sayisi'])
        ws.cell(row=row, column=3, value=float(s['toplam_miktar']))
        ws.cell(row=row, column=4, value=float(s['kalan_miktar']))
        ws.cell(row=row, column=5, value=float(s['toplam_tutar']))
        
        for col in range(1, 6):
            ws.cell(row=row, column=col).border = border
        row += 1
    
    # ===== TOPLAM SATIRI =====
    ws.cell(row=row, column=1, value="TOPLAM")
    ws.cell(row=row, column=2, value=genel_toplam['ihale_sayisi'])
    ws.cell(row=row, column=3, value=float(genel_toplam['toplam_miktar']))
    ws.cell(row=row, column=4, value=float(genel_toplam['kalan_miktar']))
    ws.cell(row=row, column=5, value=float(genel_toplam['toplam_tutar']))
    
    total_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
    for col in range(1, 6):
        c = ws.cell(row=row, column=col)
        c.font = total_font
        c.fill = total_fill
        c.border = border
    
    # ===== SÜTUN GENİŞLİKLERİ =====
    column_widths = [20, 15, 18, 18, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'urun_analiz_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

# ==================== FİİLİ STOK TAKİP RAPORU ====================
def fiili_stok_takip_raporu(request):
    """
    Fiili Stok Takip Raporu
    
    Sadece GİRİŞLER (ihale sevkleri) + ay başı devir
    """
    from django.utils.dateparse import parse_date
    from datetime import date, timedelta
    from calendar import monthrange
    from decimal import Decimal
    
    # ===== FİLTRELER =====
    donem = request.GET.get('donem', 'aylik')   # gunluk/haftalik/aylik/6aylik/yillik/ozel
    baslangic_str = request.GET.get('baslangic', '').strip()
    bitis_str = request.GET.get('bitis', '').strip()
    tedarikci_id = request.GET.get('tedarikci', '').strip()
    parti_no = request.GET.get('parti_no', '').strip()
    boy = request.GET.get('boy', '').strip()
    
    # ===== TARİH ARALIĞINI BELİRLE =====
    bugun = date.today()
    
    if donem == 'gunluk':
        baslangic = bugun
        bitis = bugun
    elif donem == 'haftalik':
        baslangic = bugun - timedelta(days=bugun.weekday())
        bitis = baslangic + timedelta(days=6)
    elif donem == 'aylik':
        baslangic = bugun.replace(day=1)
        son_gun = monthrange(bugun.year, bugun.month)[1]
        bitis = bugun.replace(day=son_gun)
    elif donem == '6aylik':
        baslangic = bugun - timedelta(days=180)
        bitis = bugun
    elif donem == 'yillik':
        baslangic = bugun.replace(month=1, day=1)
        bitis = bugun.replace(month=12, day=31)
    else:   # ozel
        baslangic = parse_date(baslangic_str) if baslangic_str else bugun.replace(day=1)
        bitis = parse_date(bitis_str) if bitis_str else bugun
    
    # ===== QUERYSET =====
    sevkler = IhaleSevk.objects.select_related(
        'ihale', 'ihale__tedarikci'
    ).filter(
        sevk_tarihi__gte=baslangic,
        sevk_tarihi__lte=bitis
    ).order_by('sevk_tarihi', 'id')
    
    # Tedarikçi filtresi
    if tedarikci_id:
        sevkler = sevkler.filter(ihale__tedarikci_id=tedarikci_id)
    
    # Parti no filtresi
    if parti_no:
        sevkler = sevkler.filter(ihale__parti_no__icontains=parti_no)
    
    # Boy filtresi
    if boy:
        sevkler = sevkler.filter(ihale__boy__icontains=boy)
    
    # ===== TOPLAMLAR =====
    toplamlar = sevkler.aggregate(
        toplam_miktar=Sum('sevk_miktar'),
        toplam_adet=Sum('sevk_adet'),
        kayit_sayisi=Count('id'),
    )
    
    # ===== DEVİR =====
    # Seçili tarih aralığının başlangıç ayı için devir
    devir = StokDevir.objects.filter(
        yil=baslangic.year,
        ay=baslangic.month
    ).first()
    
    devir_miktar = devir.devir_miktar if devir else Decimal('0')
    
    # ===== GENEL TOPLAM =====
    toplam_miktar = toplamlar['toplam_miktar'] or Decimal('0')
    genel_toplam = devir_miktar + toplam_miktar
    
    # ===== EXCEL EXPORT =====
    if request.GET.get('excel') == '1':
        return _fiili_stok_excel(
            sevkler, devir, baslangic, bitis, toplamlar, genel_toplam
        )
    
    # ===== DROPDOWN SEÇENEKLERİ =====
    tedarikciler = Tedarikci.objects.filter(
        aktif_mi=True,
        ihale__isnull=False
    ).distinct().order_by('unvan')
    
    boy_secenekleri = Ihale.objects.exclude(
        boy__isnull=True
    ).exclude(
        boy=''
    ).values_list('boy', flat=True).distinct().order_by('boy')
    
    # ===== CONTEXT =====
    context = {
        'sevkler': sevkler,
        'devir': devir,
        'devir_miktar': devir_miktar,
        'toplam_miktar': toplam_miktar,
        'toplam_adet': toplamlar['toplam_adet'] or 0,
        'genel_toplam': genel_toplam,
        'kayit_sayisi': toplamlar['kayit_sayisi'] or 0,
        'baslangic': baslangic,
        'bitis': bitis,
        'donem': donem,
        
        # Filtre değerleri
        'baslangic_str': baslangic.strftime('%Y-%m-%d'),
        'bitis_str': bitis.strftime('%Y-%m-%d'),
        'tedarikci_id': tedarikci_id,
        'parti_no': parti_no,
        'boy': boy,
        
        # Dropdown seçenekleri
        'tedarikciler': tedarikciler,
        'boy_secenekleri': boy_secenekleri,
        'donem_secenekleri': [
            ('gunluk', 'Günlük'),
            ('haftalik', 'Haftalık'),
            ('aylik', 'Aylık'),
            ('6aylik', '6 Aylık'),
            ('yillik', 'Yıllık'),
            ('ozel', 'Özel Tarih Aralığı'),
        ],
    }
    return render(request, 'stok/fiili_stok_takip_raporu.html', context)


# ==================== FİİLİ STOK EXCEL ====================
def _fiili_stok_excel(sevkler, devir, baslangic, bitis, toplamlar, genel_toplam):
    """Fiili stok takip raporu Excel export"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse
    from datetime import datetime
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Fiili Stok Takip"
    
    # ===== BAŞLIK =====
    ws.merge_cells('A1:G1')
    ws['A1'] = 'FİİLİ STOK TAKİP RAPORU'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells('A2:G2')
    ws['A2'] = f"Dönem: {baslangic.strftime('%d.%m.%Y')} - {bitis.strftime('%d.%m.%Y')}"
    ws['A2'].font = Font(bold=True, size=11)
    ws['A2'].alignment = Alignment(horizontal="center")
    
    # ===== BAŞLIKLAR =====
    headers = ['Tarih', 'Tedarikçi', 'Parti No', 'İstif No', 'Boy', 'Miktar (m³)', 'Adet']
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    
    # ===== DEVİR SATIRI =====
    row = 5
    if devir:
        ws.cell(row=row, column=1, value='DEVİR')
        ws.cell(row=row, column=2, value=f"{devir.donem_etiketi} Devri")
        ws.cell(row=row, column=3, value='-')
        ws.cell(row=row, column=4, value='-')
        ws.cell(row=row, column=5, value='-')
        ws.cell(row=row, column=6, value=float(devir.devir_miktar))
        ws.cell(row=row, column=7, value='-')
        
        # Devir satırı için renk
        for col in range(1, 8):
            c = ws.cell(row=row, column=col)
            c.border = border
            c.fill = PatternFill(start_color="FFE69C", end_color="FFE69C", fill_type="solid")
            c.font = Font(bold=True)
        row += 1
    
    # ===== SEVK SATIRLARI =====
    for s in sevkler:
        ws.cell(row=row, column=1, value=s.sevk_tarihi.strftime('%d.%m.%Y') if s.sevk_tarihi else '')
        ws.cell(row=row, column=2, value=s.ihale.tedarikci.unvan if s.ihale and s.ihale.tedarikci else '')
        ws.cell(row=row, column=3, value=s.ihale.parti_no if s.ihale else '')
        ws.cell(row=row, column=4, value=s.ihale.istif_no if s.ihale else '')
        ws.cell(row=row, column=5, value=s.ihale.boy if s.ihale else '')
        ws.cell(row=row, column=6, value=float(s.sevk_miktar or 0))
        ws.cell(row=row, column=7, value=s.sevk_adet or 0)
        
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = border
        row += 1
    
    # ===== TOPLAM SATIRLARI =====
    # Dönem toplamı
    ws.cell(row=row, column=1, value='DÖNEM TOPLAMI')
    ws.cell(row=row, column=6, value=float(toplamlar['toplam_miktar'] or 0))
    ws.cell(row=row, column=7, value=toplamlar['toplam_adet'] or 0)
    
    total_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
    for col in range(1, 8):
        c = ws.cell(row=row, column=col)
        c.font = total_font
        c.fill = total_fill
        c.border = border
    row += 1
    
    # Genel toplam
    ws.cell(row=row, column=1, value='GENEL TOPLAM (Devir + Dönem)')
    ws.cell(row=row, column=6, value=float(genel_toplam))
    
    genel_font = Font(bold=True, color="FFFFFF", size=12)
    genel_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    for col in range(1, 8):
        c = ws.cell(row=row, column=col)
        c.font = genel_font
        c.fill = genel_fill
        c.border = border
    
    # ===== SÜTUN GENİŞLİKLERİ =====
    column_widths = [15, 30, 15, 15, 20, 18, 12]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    tarih_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'fiili_stok_takip_{tarih_str}.xlsx'
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response